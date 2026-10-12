"""Export the detailed Markdown review draft to a simple, readable DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "reports" / "51_THESIS_DETAILED_REVIEW_DRAFT.md"
OUTPUT = ROOT / "reports" / "51_THESIS_DETAILED_REVIEW_FOR_COMMENTS.docx"


def set_font(run, name: str, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color: str = "D9D9D9", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, end])
    set_font(run, "Times New Roman", 9)


def add_inline_markdown(paragraph, text: str) -> None:
    """Render the small inline-Markdown subset used by the review draft."""
    text = re.sub(r"\[`([^`]+)`\]\(([^)]+)\)", r"[\1](\2)", text)
    token_pattern = re.compile(
        r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*]+\*)"
    )
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_font(run, "Times New Roman", 11)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, "Times New Roman", 11, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, "Courier New", 9.5)
        elif token.startswith("["):
            link_match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            label = link_match.group(1) if link_match else token
            run = paragraph.add_run(label)
            set_font(run, "Times New Roman", 10.5)
            run.underline = True
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, "Times New Roman", 11)
            run.italic = True
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_font(run, "Times New Roman", 11)


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    if column_count >= 5:
        headers = rows[0]
        for record_number, values in enumerate(rows[1:], start=1):
            heading = document.add_paragraph(style="Heading 3")
            heading_text = values[1] if len(values) > 1 else f"Record {record_number}"
            add_inline_markdown(heading, f"Case {record_number}  {heading_text}")
            for column_index, header in enumerate(headers):
                if column_index in (0, 1) or column_index >= len(values):
                    continue
                paragraph = document.add_paragraph()
                label_run = paragraph.add_run(f"{header}. ")
                set_font(label_run, "Times New Roman", 10.5, bold=True)
                add_inline_markdown(paragraph, values[column_index])
        return
    table = document.add_table(rows=0, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for column_index in range(column_count):
            value = row_values[column_index] if column_index < len(row_values) else ""
            cell = cells[column_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            if row_index == 0:
                shade_cell(cell, "404040")
            elif row_index % 2 == 0:
                shade_cell(cell, "F2F2F2")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_markdown(paragraph, value)
            for run in paragraph.runs:
                set_font(run, "Times New Roman", 9.5, bold=(row_index == 0))
                if row_index == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
        if row_index == 0:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.15

    style_specs = {
        "Title": (20, 0, 12),
        "Heading 1": (15, 14, 6),
        "Heading 2": (13, 11, 4),
        "Heading 3": (11.5, 9, 3),
    }
    for style_name, (size, before, after) in style_specs.items():
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    title_p_pr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)


def add_figure(document: Document, alt_text: str, raw_path: str) -> None:
    image_path = (SOURCE.parent / raw_path).resolve()
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.15))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(9)
    caption_run = caption.add_run(alt_text)
    set_font(caption_run, "Times New Roman", 9.5)
    caption_run.italic = True


def build_document() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    configure_styles(document)
    add_page_number(section.footer.paragraphs[0])

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue

        if line.startswith("!["):
            image_match = re.match(r"!\[([^\]]+)\]\(([^)]+)\)", line)
            if image_match:
                add_figure(document, image_match.group(1), image_match.group(2))
            index += 1
            continue

        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_markdown(paragraph, line[2:].strip())
            for run in paragraph.runs:
                set_font(run, "Times New Roman", 20, bold=True)
            index += 1
            continue

        if line.startswith("## "):
            heading_text = line[3:].strip()
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline_markdown(paragraph, heading_text)
            index += 1
            continue

        if line.startswith("#### "):
            paragraph = document.add_paragraph(style="Heading 3")
            add_inline_markdown(paragraph, line[5:].strip())
            index += 1
            continue

        if line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline_markdown(paragraph, line[4:].strip())
            index += 1
            continue

        if line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.7)
            paragraph.paragraph_format.right_indent = Cm(0.7)
            add_inline_markdown(paragraph, line[2:].strip())
            for run in paragraph.runs:
                run.italic = True
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and re.match(
            r"^\|?\s*:?-{3,}", lines[index + 1]
        ):
            rows = [split_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            add_table(document, rows)
            continue

        bullet_match = re.match(r"^\s*-\s+(.*)$", line)
        number_match = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if bullet_match or number_match:
            if bullet_match:
                paragraph = document.add_paragraph(style="List Bullet")
                add_inline_markdown(paragraph, bullet_match.group(1).strip())
            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.7)
                paragraph.paragraph_format.first_line_indent = Cm(-0.45)
                number_run = paragraph.add_run(f"{number_match.group(1)}.  ")
                set_font(number_run, "Times New Roman", 11)
                add_inline_markdown(paragraph, number_match.group(2).strip())
            index += 1
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.widow_control = True
        add_inline_markdown(paragraph, line)
        if line.strip() == "**Your comment:**":
            paragraph.paragraph_format.space_after = Pt(14)
        index += 1

    document.core_properties.title = "Detailed thesis review draft"
    document.core_properties.subject = "Magic Pass master thesis detailed review"
    document.core_properties.author = "Thibaud"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
